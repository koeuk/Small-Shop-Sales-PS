import csv
import hashlib
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from uuid import UUID
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from fastapi import HTTPException
from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError
from sqlalchemy import select

from .models import DataImport, Product, Sale
from .product_names import normalize_name
from .schemas import SaleInput

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10_000


def read_rows(content: bytes, filename: str):
    suffix = Path(filename).suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        raise ValueError("Only .csv and .xlsx files are supported")
    if not content:
        raise ValueError("The uploaded file is empty")
    if len(content) > MAX_FILE_BYTES:
        raise ValueError("File size must be 5 MB or less")
    if suffix == ".csv":
        try:
            rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig")), strict=True))
        except (UnicodeDecodeError, csv.Error):
            raise ValueError("CSV must be valid UTF-8 text with comma-separated columns")
    else:
        try:
            with ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 25 * 1024 * 1024:
                    raise ValueError("Expanded spreadsheet exceeds the 25 MB limit")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                sheet = workbook.active
                if (sheet.max_column or 0) > 32 or (sheet.max_row or 0) > MAX_ROWS + 1:
                    raise ValueError("Spreadsheet dimensions exceed the limit of 32 columns and 10,000 data rows. Remove unused rows/columns and save again.")
                sheet.reset_dimensions()
                rows = []
                for row in sheet.iter_rows(values_only=True, max_col=33):
                    if row[32] is not None:
                        raise ValueError("Spreadsheet may contain at most 32 columns")
                    values = list(row[:32])
                    while values and values[-1] is None:
                        values.pop()
                    rows.append(values)
                    if len(rows) > MAX_ROWS + 1:
                        raise ValueError(f"Import is limited to {MAX_ROWS:,} rows")
            finally:
                workbook.close()
        except (BadZipFile, KeyError, OSError, ParseError, InvalidFileException, DefusedXmlException):
            raise ValueError("Unable to read this XLSX workbook")
    rows = [row for row in rows if any(value is not None and str(value).strip() for value in row)]
    if len(rows) < 2:
        raise ValueError("Include a header and at least one sales row")
    if len(rows) - 1 > MAX_ROWS:
        raise ValueError(f"Import is limited to {MAX_ROWS:,} rows")
    headers = [str(value or "").strip().lower() for value in rows[0]]
    if len(headers) != len(set(headers)) or "" in headers:
        raise ValueError("Column headers must be unique and nonempty")
    required = {"sale_date", "quantity", "unit_price"}
    if not required.issubset(headers) or not {"product_id", "product_name"}.intersection(headers):
        raise ValueError("Required columns: sale_date, product_name (or product_id), quantity, unit_price")
    return headers, rows[1:]


def prepare_import(session, owner_id, content, filename):
    fingerprint = hashlib.sha256(content).hexdigest()
    existing = session.scalar(select(DataImport).where(DataImport.owner_id == owner_id, DataImport.fingerprint == fingerprint))
    if existing and existing.status == "completed":
        raise HTTPException(409, "This file has already been imported. Rename does not bypass duplicate detection.")
    products = list(session.scalars(select(Product).where(Product.owner_id == owner_id, Product.archived.is_(False))))
    by_id = {product.id: product for product in products}
    by_name = {}
    for product in products:
        by_name.setdefault(normalize_name(product.name), []).append(product)
    valid, errors, preview = [], [], []
    try:
        headers, rows = read_rows(content, filename)
    except (ValueError, TypeError, OverflowError) as exc:
        return {"fingerprint": fingerprint, "total_rows": 0, "valid_rows": 0, "invalid_rows": 0, "errors": [{"row": 0, "message": str(exc)}], "preview": [], "fatal": True}, [], existing
    for number, values in enumerate(rows, start=2):
        row = dict(zip(headers, values))
        error = None
        try:
            if len(values) > len(headers):
                raise ValueError("More values than column headers")
            product_id = str(row.get("product_id") or "").strip()
            product_name = str(row.get("product_name") or "").strip()
            matches = by_name.get(normalize_name(product_name), [])
            if not product_id and len(matches) > 1:
                raise ValueError("Product name is ambiguous; use product_id or rename duplicate products")
            product = by_id.get(str(UUID(product_id))) if product_id else next(iter(matches), None)
            if not product:
                raise ValueError("Product is unknown or archived; create or restore it before importing")
            if product_id and product_name and normalize_name(product.name) != normalize_name(product_name):
                raise ValueError("product_id and product_name refer to different products")
            quantity = str(row.get("quantity", "")).strip()
            if not re.fullmatch(r"[1-9]\d*", quantity):
                raise ValueError("Quantity must be a positive whole number")
            day = row.get("sale_date")
            if isinstance(day, datetime):
                day = day.date()
            elif not isinstance(day, date):
                day = str(day).strip()
                if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", day):
                    raise ValueError("Dates must use YYYY-MM-DD")
                day = date.fromisoformat(day)
            sale = SaleInput(product_id=product.id, sale_date=day, quantity=int(quantity), unit_price=Decimal(str(row.get("unit_price", "")).strip()))
            valid.append({**sale.model_dump(), "product_id": product.id, "total_amount": sale.quantity * sale.unit_price, "source_row": number})
        except (ValidationError, ValueError, TypeError, InvalidOperation, OverflowError) as exc:
            error = "; ".join(item["msg"] for item in exc.errors()) if isinstance(exc, ValidationError) else str(exc)
            errors.append({"row": number, "message": error})
        if len(preview) < 15:
            preview.append({"row": number, "sale_date": str(row.get("sale_date", "")), "product": str(row.get("product_name") or row.get("product_id") or ""), "quantity": str(row.get("quantity", "")), "unit_price": str(row.get("unit_price", "")), "error": error})
    return {"fingerprint": fingerprint, "total_rows": len(rows), "valid_rows": len(valid), "invalid_rows": len(errors), "errors": errors[:100], "preview": preview, "fatal": False}, valid, existing


def commit_import(session, owner_id, filename, summary, valid, existing, synthetic):
    record = existing or DataImport(owner_id=owner_id, fingerprint=summary["fingerprint"], file_name=filename[:255], file_type=Path(filename).suffix.lstrip(".").lower()[:20])
    record.total_rows = summary["total_rows"]
    record.row_count = len(valid)
    record.invalid_rows = summary["invalid_rows"]
    record.status = "completed" if valid else "failed"
    record.error_summary = "\n".join(f"Row {item['row']}: {item['message']}" for item in summary["errors"]) or None
    record.synthetic = synthetic
    session.add(record)
    session.flush()
    for values in valid:
        session.add(Sale(owner_id=owner_id, import_id=record.id, **values))
    session.flush()
    return record
